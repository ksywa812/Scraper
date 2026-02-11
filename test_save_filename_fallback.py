import os
from openpyxl import load_workbook
from scraper import save_to_excel


def test_save_invalid_filename_fallback(monkeypatch, tmp_path):
    # Run inside temporary working directory so 'scraped' is created under tmp_path
    monkeypatch.chdir(tmp_path)

    data = [{
        'name': 'Fallback Test Salon',
        'formatted_address': 'Test Address',
        'formatted_phone_number': '000',
        'emails': ['one@test.com']
    }]

    # Provide an obviously invalid filename that previously caused attempts to save to root ("\\.xlsx")
    invalid_name = "\\.xlsx"

    # Should not raise and should create a valid file under tmp_path/scraped/
    save_to_excel(data, filename=invalid_name, city='TestCity', append=False)

    fallback_path = tmp_path / 'scraped' / 'results.xlsx'
    assert fallback_path.exists(), f"Expected fallback file at {fallback_path}"

    wb = load_workbook(str(fallback_path))
    ws = wb.active

    # There should be at least one header row and one data row (so max_row >= 2)
    assert ws.max_row >= 2
