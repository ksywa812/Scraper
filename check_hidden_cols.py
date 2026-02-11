from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from pathlib import Path
p = Path('scraped') / 'IdeaMusicLeads.xlsx'
wb = load_workbook(p)
s = wb[wb.sheetnames[0]]
headers = [cell.value for cell in next(s.iter_rows(min_row=1, max_row=1))]
print('Headers:', headers)
for col_name in ('Adres E-mail 2','Adres E-mail 3'):
    try:
        idx = headers.index(col_name) + 1
        letter = get_column_letter(idx)
        hidden = s.column_dimensions[letter].hidden
        width = s.column_dimensions[letter].width
        print(col_name, 'col', letter, 'hidden=', hidden, 'width=', width)
    except Exception as e:
        print('Error for', col_name, e)
