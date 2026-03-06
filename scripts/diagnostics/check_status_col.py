from pathlib import Path

from openpyxl import load_workbook

PROJECT_ROOT = Path(__file__).resolve().parents[2]
p = PROJECT_ROOT / "Data" / "Raport" / "IdeaMusicLeads.xlsx"
wb = load_workbook(p)
s = wb[wb.sheetnames[0]]
headers = [cell.value for cell in next(s.iter_rows(min_row=1, max_row=1))]
status_idx = headers.index("Status") + 1
print("Status column index (1-based):", status_idx)
for row in s.iter_rows(min_row=s.max_row - 10, max_row=s.max_row, values_only=True):
    print("Status value:", row[status_idx - 1])
