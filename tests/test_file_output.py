"""Tests for file output functions: save_to_csv, save_to_json, save_to_excel."""
import csv
import json
import sys
import os

sys.path.insert(0, os.path.join(os.path.dirname(__file__), ".."))

import pytest
from scraper import save_to_csv, save_to_json, save_to_excel


# ---------------------------------------------------------------------------
# Sample data
# ---------------------------------------------------------------------------

SAMPLE_DATA = [
    {
        "name": "Salon Wellness",
        "formatted_address": "ul. Kwiatowa 1, Wrocław",
        "formatted_phone_number": "+48 71 123 45 67",
        "website": "https://salon-wellness.pl",
        "emails": ["kontakt@salon-wellness.pl", "biuro@salon-wellness.pl"],
        "sources": ["panorama", "pkt"],
    },
    {
        "name": "Spa Relax",
        "formatted_address": "ul. Różana 3, Kraków",
        "formatted_phone_number": "+48 12 345 67 89",
        "website": "",
        "emails": [],
        "sources": ["krs"],
    },
]


# ---------------------------------------------------------------------------
# save_to_csv
# ---------------------------------------------------------------------------

class TestSaveToCsv:
    def test_creates_file(self, tmp_path):
        path = str(tmp_path / "output.csv")
        save_to_csv(SAMPLE_DATA, path)
        assert os.path.exists(path)

    def test_csv_has_header_row(self, tmp_path):
        path = str(tmp_path / "output.csv")
        save_to_csv(SAMPLE_DATA, path)
        with open(path, newline="", encoding="utf-8") as f:
            reader = csv.reader(f)
            header = next(reader)
        assert "Name" in header
        assert "Email 1" in header

    def test_csv_has_correct_row_count(self, tmp_path):
        path = str(tmp_path / "output.csv")
        save_to_csv(SAMPLE_DATA, path)
        with open(path, newline="", encoding="utf-8") as f:
            rows = list(csv.reader(f))
        # header + 2 data rows
        assert len(rows) == 3

    def test_csv_name_written(self, tmp_path):
        path = str(tmp_path / "output.csv")
        save_to_csv(SAMPLE_DATA, path)
        with open(path, newline="", encoding="utf-8") as f:
            rows = list(csv.reader(f))
        names = [r[0] for r in rows[1:]]
        assert "Salon Wellness" in names

    def test_csv_email1_written(self, tmp_path):
        path = str(tmp_path / "output.csv")
        save_to_csv(SAMPLE_DATA, path)
        with open(path, newline="", encoding="utf-8") as f:
            rows = list(csv.reader(f))
        # First data row, Email 1 is column index 4
        assert "kontakt@salon-wellness.pl" in rows[1]

    def test_csv_sources_written(self, tmp_path):
        path = str(tmp_path / "output.csv")
        save_to_csv(SAMPLE_DATA, path)
        with open(path, newline="", encoding="utf-8") as f:
            rows = list(csv.reader(f))
        # Sources is the last column
        assert "panorama" in rows[1][-1]


# ---------------------------------------------------------------------------
# save_to_json
# ---------------------------------------------------------------------------

class TestSaveToJson:
    def test_creates_file(self, tmp_path):
        path = str(tmp_path / "output.json")
        save_to_json(SAMPLE_DATA, path)
        assert os.path.exists(path)

    def test_json_is_valid(self, tmp_path):
        path = str(tmp_path / "output.json")
        save_to_json(SAMPLE_DATA, path)
        with open(path, encoding="utf-8") as f:
            data = json.load(f)
        assert isinstance(data, list)

    def test_json_has_correct_length(self, tmp_path):
        path = str(tmp_path / "output.json")
        save_to_json(SAMPLE_DATA, path)
        with open(path, encoding="utf-8") as f:
            data = json.load(f)
        assert len(data) == 2

    def test_json_name_preserved(self, tmp_path):
        path = str(tmp_path / "output.json")
        save_to_json(SAMPLE_DATA, path)
        with open(path, encoding="utf-8") as f:
            data = json.load(f)
        assert data[0]["name"] == "Salon Wellness"

    def test_json_emails_preserved(self, tmp_path):
        path = str(tmp_path / "output.json")
        save_to_json(SAMPLE_DATA, path)
        with open(path, encoding="utf-8") as f:
            data = json.load(f)
        assert "kontakt@salon-wellness.pl" in data[0]["emails"]

    def test_json_unicode_preserved(self, tmp_path):
        path = str(tmp_path / "output.json")
        save_to_json(SAMPLE_DATA, path)
        with open(path, encoding="utf-8") as f:
            content = f.read()
        assert "Wrocław" in content


# ---------------------------------------------------------------------------
# save_to_excel (basic — no template)
# ---------------------------------------------------------------------------

class TestSaveToExcel:
    def test_creates_xlsx_file(self, tmp_path):
        path = str(tmp_path / "output.xlsx")
        save_to_excel(SAMPLE_DATA, filename=path)
        assert os.path.exists(path)

    def test_xlsx_readable(self, tmp_path):
        import openpyxl
        path = str(tmp_path / "output.xlsx")
        save_to_excel(SAMPLE_DATA, filename=path)
        wb = openpyxl.load_workbook(path)
        assert wb is not None

    def test_xlsx_has_data_rows(self, tmp_path):
        import openpyxl
        path = str(tmp_path / "output.xlsx")
        save_to_excel(SAMPLE_DATA, filename=path)
        wb = openpyxl.load_workbook(path)
        ws = wb.active
        # At least header + 2 data rows
        assert ws.max_row >= 3
