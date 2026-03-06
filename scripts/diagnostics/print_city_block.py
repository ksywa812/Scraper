import sys
from pathlib import Path

from openpyxl import load_workbook

PROJECT_ROOT = Path(__file__).resolve().parents[2]

city = sys.argv[1] if len(sys.argv) > 1 else "Bydgoszcz"
wb = load_workbook(PROJECT_ROOT / "Data" / "Raport" / "IdeaMusicLeads.xlsx")
ws = wb[wb.sheetnames[0]]
rows = list(ws.iter_rows(values_only=True))
city_up = city.upper()
matches = []
for i, row in enumerate(rows, start=1):
    v = row[0]
    if not v:
        continue
    s = str(v).strip().upper()
    if s == city_up or s == f"!{city_up}":
        matches.append(i)

if not matches:
    print(f"No header for {city} found.")
else:
    last_header = matches[-1]
    print(f"Found header for {city} at row {last_header}.")
    next_header = None
    for i in range(last_header + 1, len(rows) + 1):
        v = rows[i - 1][0]
        if v:
            s = str(v).strip().upper()
            if len(s.split()) <= 4 and s == s.upper() and (s.startswith("!") or s == s):
                if ("," not in s) and (not any(c.isdigit() for c in s)):
                    next_header = i
                    break
    if next_header:
        print(
            f"Next header at row {next_header}. Rows for {city} are {last_header + 1}..{next_header - 1}"
        )
    else:
        print(f"No next header. Rows for {city} are {last_header + 1}..{len(rows)}")
    for r in range(last_header, min(last_header + 10, len(rows))):
        print(r + 1, rows[r])
