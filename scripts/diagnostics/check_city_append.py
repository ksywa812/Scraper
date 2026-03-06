from pathlib import Path

from openpyxl import load_workbook

PROJECT_ROOT = Path(__file__).resolve().parents[2]
p = PROJECT_ROOT / "Data" / "Raport" / "IdeaMusicLeads.xlsx"
wb = load_workbook(p)
s = wb[wb.sheetnames[0]]
rows = list(s.iter_rows(values_only=False))

n = len(rows)
print("Total rows:", n)

header = [c.value for c in rows[0]]
try:
    status_idx = header.index("Status")
except ValueError:
    status_idx = None

print("Header Index of Status:", status_idx)
for i, row in enumerate(rows[-30:], start=n - 30 + 1):
    vals = [cell.value for cell in row[:10]]
    print(i, vals)
    if row[0].fill and row[0].fill.fill_type:
        if row[0].fill.fill_type == "solid":
            print("  -> Row has fill color:", row[0].fill.start_color.index)
