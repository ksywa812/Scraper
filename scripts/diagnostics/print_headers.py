import sys
from pathlib import Path

from openpyxl import load_workbook

PROJECT_ROOT = Path(__file__).resolve().parents[2]

fname = sys.argv[1] if len(sys.argv) > 1 else "test_spa_wroclaw2.xlsx"
p = PROJECT_ROOT / "Data" / "Raport" / fname
wb = load_workbook(p, read_only=True)
print("Sheets:", wb.sheetnames)
s = wb[wb.sheetnames[0]]
headers = [cell.value for cell in next(s.iter_rows(min_row=1, max_row=1))]
print("Headers:", headers)
rows = list(s.iter_rows(min_row=1, max_row=3, values_only=True))
for r in rows:
    print(r)
